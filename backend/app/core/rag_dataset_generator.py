from __future__ import annotations

import asyncio
import hashlib
import json
import logging
import math
import os
import re
import shutil
import uuid
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Iterable

from sqlalchemy.orm import Session, joinedload

from app.core.config import settings
from app.core.prompt_manager import (
    QUESTION_GENERATION_PROMPT,
    QUESTION_VALIDATION_PROMPT,
)
from app.models.dataset import Dataset, DatasetRow
from app.models.llm_config import LLMConfig
from app.models.rag_dataset_job import (
    RagDatasetChunk,
    RagDatasetDocument,
    RagDatasetJob,
    RagDatasetSample,
)

logger = logging.getLogger(__name__)

RAG_JOB_TIMEOUT_SECONDS = 15 * 60
RAG_QUESTION_GENERATION_TIMEOUT_SECONDS = 120

_PAGE_LINE_RE = re.compile(r"^\s*第?\s*\d+\s*页(?:\s*/\s*共?\s*\d+\s*页)?\s*$")
_DATE_RE = re.compile(r"\d{4}\s*年\s*\d{1,2}\s*月\s*\d{1,2}\s*日")
_SECTION_NUMBER_RE = re.compile(r"^\d+(?:\.\d+){1,4}$")
_REVISION_META_KEYWORDS = (
    "编制",
    "文件编号",
    "版本号",
    "更改页码",
    "更改条款号",
    "新建/修订",
    "审核人",
    "批准",
    "修订日期",
    "备注",
)
_FORM_SECTION_MARKERS = (
    "回执",
    "回条",
    "家长签名",
    "学生姓名",
    "签字",
    "签名",
    "报名缴费",
    "由此进入",
    "班主任",
    "不需要据下",
    "我已认真阅读",
    "我的孩子将参加",
    "原因不能参加",
    "在家学习",
    "到校学习",
    "安全管理",
)
_OCR_GARBLED_RE = re.compile(r"[A-Za-z]{2,}\s*[A-Za-z]{2,}")
_STRUCTURED_HEADING_RE = re.compile(r"^\s*(\d+(?:\.\d+){1,5})(?:[\.、．])?\s*(.+?)\s*$")
_STRUCTURED_HEADING_NUMBER_RE = re.compile(r"^\s*(\d+(?:\.\d+){1,5})\s*$")


def _normalize_pdf_line(line: str) -> str:
    return re.sub(r"\s+", " ", (line or "").strip())


def _looks_like_header_footer(text: str) -> bool:
    normalized = _normalize_pdf_line(text)
    compact = re.sub(r"\s+", "", normalized)
    if not normalized:
        return True
    if _PAGE_LINE_RE.match(normalized):
        return True
    if "第" in normalized and "页" in normalized and _DATE_RE.search(normalized):
        return True
    if _SECTION_NUMBER_RE.match(compact):
        return True
    if _DATE_RE.search(normalized) and len(normalized) <= 40:
        return True
    if len(compact) <= 12 and not re.search(r"[\u4e00-\u9fffA-Za-z]{4,}", compact):
        return True
    return False


def _looks_like_revision_metadata(text: str) -> bool:
    normalized = _normalize_pdf_line(text)
    keyword_hits = sum(1 for keyword in _REVISION_META_KEYWORDS if keyword in normalized)
    if len(normalized) <= 180 and keyword_hits >= 2:
        return True
    if len(normalized) <= 180 and normalized.startswith("注：1.") and "修订" in normalized:
        return True
    if len(normalized) <= 180 and normalized.startswith("注：") and ("修订" in normalized or "审核" in normalized or "批准" in normalized):
        return True
    if len(normalized) <= 180 and re.match(r"^\d+\.", normalized) and ("修订" in normalized or "审核" in normalized or "批准" in normalized):
        return True
    return False


def _looks_like_form_marker(text: str) -> bool:
    normalized = _normalize_pdf_line(text)
    return any(marker in normalized for marker in _FORM_SECTION_MARKERS)


def _looks_like_ocr_garble(text: str) -> bool:
    normalized = _normalize_pdf_line(text)
    if not normalized:
        return True
    compact = re.sub(r"\s+", "", normalized)
    cjk_chars = len(re.findall(r"[\u4e00-\u9fff]", compact))
    ascii_chars = len(re.findall(r"[A-Za-z]", compact))
    digits = len(re.findall(r"\d", compact))
    punct = len(re.findall(r"[^A-Za-z0-9\u4e00-\u9fff]", compact))
    if ascii_chars >= 6 and cjk_chars <= 4 and digits <= 2:
        return True
    if _OCR_GARBLED_RE.search(normalized) and cjk_chars <= 6:
        return True
    if punct >= max(4, len(compact) * 0.25) and cjk_chars <= 6:
        return True
    return False


def _prune_form_tail(lines: list[str]) -> list[str]:
    if not lines:
        return lines
    marker_index: int | None = None
    for idx, line in enumerate(lines):
        if idx < max(3, int(len(lines) * 0.4)):
            continue
        if _looks_like_form_marker(line):
            marker_index = idx
            break
    if marker_index is None:
        return lines
    return lines[:marker_index]


def evaluate_chunk_quality(text: str) -> dict[str, Any]:
    normalized_text = _normalize_pdf_line(text)
    compact = re.sub(r"\s+", "", normalized_text)
    cjk_chars = len(re.findall(r"[\u4e00-\u9fff]", compact))
    word_chars = len(re.findall(r"[A-Za-z]", compact))
    digits = len(re.findall(r"\d", compact))
    sentence_markers = len(re.findall(r"[。！？.!?；;:：]", text))
    reasons: list[str] = []
    score = 1.0

    if _looks_like_header_footer(normalized_text):
        return {"score": 0.0, "label": "filtered", "reasons": ["页眉页脚/页码噪声"]}
    if _looks_like_revision_metadata(normalized_text):
        return {"score": 0.05, "label": "filtered", "reasons": ["制度修订记录或版本元信息"]}
    if _looks_like_form_marker(normalized_text):
        return {"score": 0.05, "label": "filtered", "reasons": ["表单/回执字段"]}
    if _looks_like_ocr_garble(normalized_text):
        return {"score": 0.15, "label": "filtered", "reasons": ["OCR 乱码或低可读文本"]}

    if len(compact) < 40:
        score -= 0.45
        reasons.append("文本过短")
    elif len(compact) < 90:
        score -= 0.15
        reasons.append("信息量偏少")

    if cjk_chars < 12 and word_chars < 24:
        score -= 0.25
        reasons.append("正文特征较弱")

    if digits >= max(8, int(len(compact) * 0.28)) and sentence_markers == 0:
        score -= 0.2
        reasons.append("编号/数字占比偏高")

    if sentence_markers == 0 and len(compact) < 140:
        score -= 0.15
        reasons.append("缺少完整句子")

    if len(compact) < 220 and re.search(r"(因|为|及|等|并|或|由|按)$", compact):
        score -= 0.2
        reasons.append("疑似半句尾段")

    if re.search(r"(步骤|流程|注意事项|条件|规则|要求|说明|如下|适用|负责|应当|可以|不得)", normalized_text):
        score += 0.12
        reasons.append("包含规则/流程型正文")

    if len(compact) >= 220:
        score += 0.08
        reasons.append("正文信息较充足")

    score = max(0.0, min(1.0, round(score, 2)))
    if score >= 0.75:
        label = "good"
    elif score >= 0.45:
        label = "medium"
    elif score >= 0.2:
        label = "low"
    else:
        label = "filtered"

    return {"score": score, "label": label, "reasons": reasons or ["普通正文"]}


def _is_valid_structured_heading(number: str, title: str) -> bool:
    title_text = _normalize_pdf_line(title)
    compact = re.sub(r"\s+", "", title_text)
    if not compact:
        return False
    if _looks_like_revision_metadata(f"{number} {title_text}"):
        return False
    if _looks_like_form_marker(title_text):
        return False
    if len(compact) <= 2 and not re.search(r"[\u4e00-\u9fffA-Za-z]{2,}", compact):
        return False
    return True


def _make_section_node(number: str, heading_line: str) -> dict[str, Any]:
    return {
        "number": number,
        "heading_line": heading_line.strip(),
        "body_lines": [],
        "children": [],
    }


def _parse_structured_sections(text: str) -> dict[str, Any]:
    root = _make_section_node("", "")
    current = root
    stack: list[tuple[int, dict[str, Any]]] = [(0, root)]
    pending_heading_number: str | None = None

    for raw_line in text.split("\n"):
        line = _normalize_pdf_line(raw_line)
        if not line:
            continue

        only_number = _STRUCTURED_HEADING_NUMBER_RE.match(line)
        if only_number:
            pending_heading_number = only_number.group(1)
            continue

        if pending_heading_number:
            if _is_valid_structured_heading(pending_heading_number, line):
                number = pending_heading_number
                level = number.count(".") + 1
                heading_line = f"{number} {line}".strip()
                node = _make_section_node(number, heading_line)
                while stack and stack[-1][0] >= level:
                    stack.pop()
                parent = stack[-1][1] if stack else root
                parent["children"].append(node)
                stack.append((level, node))
                current = node
                pending_heading_number = None
                continue
            current["body_lines"].append(pending_heading_number)
            pending_heading_number = None

        matched = _STRUCTURED_HEADING_RE.match(line)
        if matched and _is_valid_structured_heading(matched.group(1), matched.group(2)):
            number = matched.group(1)
            title = matched.group(2).strip()
            level = number.count(".") + 1
            node = _make_section_node(number, f"{number} {title}".strip())
            while stack and stack[-1][0] >= level:
                stack.pop()
            parent = stack[-1][1] if stack else root
            parent["children"].append(node)
            stack.append((level, node))
            current = node
            continue

        current["body_lines"].append(line)

    if pending_heading_number:
        current["body_lines"].append(pending_heading_number)

    return root


def _node_text(node: dict[str, Any]) -> str:
    lines: list[str] = []
    if node.get("heading_line"):
        lines.append(node["heading_line"])
    lines.extend(node.get("body_lines") or [])
    return "\n".join(line for line in lines if line).strip()


def _node_is_substantive(node: dict[str, Any]) -> bool:
    text = _node_text(node)
    compact = re.sub(r"\s+", "", text)
    return bool(text) and len(compact) >= 24 and not _is_low_information_text(text)


def _collect_structured_section_entries(
    node: dict[str, Any],
    context_headings: tuple[str, ...],
    entries: list[dict[str, Any]],
) -> None:
    node_text = _node_text(node)
    if node.get("number"):
        parent_key = node["number"].rsplit(".", 1)[0] if "." in node["number"] else node["number"]
    else:
        parent_key = "intro"

    if node.get("heading_line") and not node.get("children"):
        lines = [*context_headings[-2:], node_text]
        entries.append(
            {
                "number": node.get("number") or "intro",
                "parent_key": parent_key,
                "text": "\n".join(line for line in lines if line).strip(),
            }
        )
        return

    next_context = context_headings
    if node.get("heading_line"):
        next_context = (*context_headings, node["heading_line"])
        if _node_is_substantive(node):
            lines = [*context_headings[-2:], node_text]
            entries.append(
                {
                    "number": node.get("number") or "intro",
                    "parent_key": parent_key,
                    "text": "\n".join(line for line in lines if line).strip(),
                }
            )

    for child in node.get("children") or []:
        _collect_structured_section_entries(child, next_context, entries)


def _length_chunk_text(text: str, chunk_size: int, overlap: int) -> list[str]:
    cleaned = text.replace("\r\n", "\n").strip()
    if not cleaned:
        return []
    if len(cleaned) <= chunk_size:
        return [cleaned]

    paragraphs = [part.strip() for part in cleaned.split("\n") if part.strip()]
    chunks: list[str] = []
    current = ""

    def _flush() -> None:
        nonlocal current
        if current.strip():
            chunks.append(current.strip())
        current = ""

    for paragraph in paragraphs:
        candidate = f"{current}\n{paragraph}".strip() if current else paragraph
        if len(candidate) <= chunk_size:
            current = candidate
            continue

        if current:
            _flush()

        remaining = paragraph
        while len(remaining) > chunk_size:
            cut = chunk_size
            for sep in ["\n", "。", "；", ".", ";", "，", ",", " "]:
                pos = remaining[:chunk_size].rfind(sep)
                if pos > int(chunk_size * 0.45):
                    cut = pos + 1
                    break
            piece = remaining[:cut].strip()
            if piece:
                chunks.append(piece)
            remaining = remaining[max(0, cut - overlap):].strip()
        current = remaining

    _flush()
    return [chunk for chunk in chunks if chunk]


def _merge_structured_entry_text(current_text: str, next_text: str) -> str:
    current_lines = [line.strip() for line in current_text.split("\n") if line.strip()]
    next_lines = [line.strip() for line in next_text.split("\n") if line.strip()]
    if not current_lines:
        return "\n".join(next_lines).strip()
    if not next_lines:
        return "\n".join(current_lines).strip()

    merged_lines = list(current_lines)
    start_index = 0
    if (
        next_lines
        and _STRUCTURED_HEADING_RE.match(next_lines[0])
        and next_lines[0] in merged_lines
    ):
        start_index = 1
    while (
        start_index < len(next_lines)
        and start_index < len(current_lines)
        and next_lines[start_index] == current_lines[start_index]
    ):
        start_index += 1
    merged_lines.extend(next_lines[start_index:])
    return "\n".join(merged_lines).strip()


def _split_structured_chunks(text: str, chunk_size: int, overlap: int) -> list[str]:
    root = _parse_structured_sections(text)
    entries: list[dict[str, Any]] = []

    intro_text = "\n".join(root.get("body_lines") or []).strip()
    if intro_text and not _is_low_information_text(intro_text):
        entries.append({"number": "intro", "parent_key": "intro", "text": intro_text})

    for child in root.get("children") or []:
        _collect_structured_section_entries(child, (), entries)

    entries = [entry for entry in entries if entry.get("text")]
    if len(entries) < 2:
        return []

    if (
        len(entries) >= 2
        and entries[0].get("number") == "intro"
        and len(re.sub(r"\s+", "", entries[0]["text"])) < max(80, int(chunk_size * 0.18))
    ):
        entries[1]["text"] = f"{entries[0]['text'].rstrip()}\n{entries[1]['text'].lstrip()}".strip()
        entries = entries[1:]

    merged: list[str] = []
    current = entries[0].copy()
    for entry in entries[1:]:
        current_len = len(current["text"])
        next_len = len(entry["text"])
        same_parent = current.get("parent_key") == entry.get("parent_key")
        should_merge = (
            same_parent
            and (
                current_len < max(320, int(chunk_size * 0.38))
                or next_len < max(180, int(chunk_size * 0.22))
            )
            and current_len + next_len <= int(chunk_size * 1.35)
        )
        if should_merge:
            current["text"] = _merge_structured_entry_text(current["text"], entry["text"])
        else:
            merged.append(current["text"])
            current = entry.copy()
    merged.append(current["text"])

    final_chunks: list[str] = []
    for entry_text in merged:
        if len(entry_text) <= int(chunk_size * 1.2):
            final_chunks.append(entry_text.strip())
        else:
            final_chunks.extend(_length_chunk_text(entry_text, chunk_size, overlap))
    return [chunk for chunk in final_chunks if chunk]


def _is_pdf_noise_line(line: str, repeated_lines: set[str]) -> bool:
    normalized = _normalize_pdf_line(line)
    if not normalized:
        return True
    if _looks_like_header_footer(normalized):
        return True
    if _looks_like_revision_metadata(normalized):
        return True
    if normalized in repeated_lines and len(normalized) <= 120:
        return True
    compact = re.sub(r"\s+", "", normalized)
    if len(compact) <= 8 and not re.search(r"[\u4e00-\u9fffA-Za-z]{4,}", compact):
        return True
    return False


def _extract_pdf_with_fitz(content: bytes) -> str:
    try:
        import fitz  # type: ignore
    except Exception as exc:  # pragma: no cover - optional dependency
        raise ValueError("当前环境未安装 PyMuPDF 解析依赖，无法执行版面感知 PDF 提取") from exc

    doc = fitz.open(stream=content, filetype="pdf")
    try:
        repeated_blocks: dict[str, int] = {}
        page_blocks: list[list[tuple[str, float]]] = []
        for page in doc:
            page_height = float(page.rect.height or 0)
            top_band = page_height * 0.12
            bottom_band = page_height * 0.88
            raw_blocks = page.get_text("blocks", sort=True) or []
            kept_blocks: list[tuple[str, float]] = []
            for block in raw_blocks:
                if len(block) < 5:
                    continue
                x0, y0, x1, y1, text = block[:5]
                normalized = _normalize_pdf_line(str(text))
                if not normalized:
                    continue
                block_height = max(1.0, float(y1) - float(y0))
                block_center = (float(y0) + float(y1)) / 2
                in_margin = block_center <= top_band or block_center >= bottom_band
                if in_margin and _looks_like_header_footer(normalized):
                    continue
                if _looks_like_header_footer(normalized) and len(normalized) <= 100:
                    continue
                density = len(re.sub(r"\s+", "", normalized)) / max(1.0, float(x1) - float(x0))
                if density < 0.015 and len(normalized) < 24:
                    continue
                kept_blocks.append((normalized, block_height))
                repeated_blocks[normalized] = repeated_blocks.get(normalized, 0) + 1
            page_blocks.append(kept_blocks)

        repeated_texts = {
            text
            for text, count in repeated_blocks.items()
            if count >= 2 and (count / max(1, len(doc))) >= 0.5 and len(text) <= 120
        }

        cleaned_pages: list[str] = []
        for blocks in page_blocks:
            page_lines: list[str] = []
            for normalized, _block_height in blocks:
                if normalized in repeated_texts and _looks_like_header_footer(normalized):
                    continue
                if _is_low_information_text(normalized):
                    continue
                page_lines.append(normalized)
            page_lines = _prune_form_tail(page_lines)
            page_text = "\n".join(page_lines).strip()
            if page_text:
                cleaned_pages.append(page_text)

        return "\n\n".join(cleaned_pages).strip()
    finally:
        doc.close()


def _ocr_pdf_with_fitz(content: bytes) -> str:
    try:
        import fitz  # type: ignore
    except Exception as exc:  # pragma: no cover - optional dependency
        raise ValueError("当前环境未安装 PyMuPDF，无法执行 OCR") from exc

    if shutil.which("tesseract") is None:
        raise ValueError("当前环境未安装 tesseract，无法对扫描 PDF 执行 OCR")

    doc = fitz.open(stream=content, filetype="pdf")
    try:
        page_texts: list[str] = []
        for page in doc:
            try:
                textpage = page.get_textpage_ocr(
                    flags=0,
                    language="chi_sim+eng",
                    dpi=300,
                    full=True,
                )
                text = page.get_text("text", textpage=textpage)
            except Exception:
                text = ""
            if text.strip():
                page_texts.append(text)
        return _clean_pdf_text(page_texts)
    finally:
        doc.close()


def _clean_pdf_text(page_texts: list[str]) -> str:
    if not page_texts:
        return ""

    per_page_lines: list[list[str]] = []
    frequency: dict[str, int] = {}
    for text in page_texts:
        lines = [_normalize_pdf_line(line) for line in (text or "").splitlines()]
        lines = [line for line in lines if line]
        per_page_lines.append(lines)
        for line in set(lines):
            frequency[line] = frequency.get(line, 0) + 1

    repeated_lines = {
        line
        for line, count in frequency.items()
        if count >= 2 and (count / max(1, len(page_texts))) >= 0.5
    }

    cleaned_pages: list[str] = []
    for lines in per_page_lines:
        filtered = [
            line
            for line in lines
            if not _is_pdf_noise_line(line, repeated_lines)
            and not _looks_like_form_marker(line)
            and not _looks_like_ocr_garble(line)
        ]
        filtered = _prune_form_tail(filtered)
        page_text = "\n".join(filtered).strip()
        if page_text:
            cleaned_pages.append(page_text)
    return "\n\n".join(cleaned_pages).strip()


def _is_low_information_text(text: str) -> bool:
    quality = evaluate_chunk_quality(text)
    return quality["score"] < 0.35


def parse_uploaded_file(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix in {".txt", ".md", ".markdown", ".csv", ".json"}:
        for encoding in ("utf-8", "utf-8-sig", "gbk"):
            try:
                return content.decode(encoding)
            except UnicodeDecodeError:
                continue
        return content.decode("utf-8", errors="ignore")

    if suffix == ".docx":
        try:
            from docx import Document  # type: ignore
        except Exception as exc:  # pragma: no cover - optional dependency
            raise ValueError("当前环境未安装 docx 解析依赖，无法读取 .docx 文件") from exc
        from io import BytesIO

        doc = Document(BytesIO(content))
        return "\n".join(paragraph.text for paragraph in doc.paragraphs if paragraph.text)

    if suffix == ".pdf":
        fitz_exc: Exception | None = None
        try:
            cleaned = _extract_pdf_with_fitz(content)
            if cleaned:
                return cleaned
        except Exception as exc:  # pragma: no cover - fallback path
            fitz_exc = exc

        try:
            cleaned = _ocr_pdf_with_fitz(content)
            if cleaned:
                return cleaned
        except Exception:
            pass

        try:
            import pdfplumber  # type: ignore
        except Exception as exc:  # pragma: no cover - optional dependency
            if fitz_exc is not None:
                raise ValueError("当前环境未安装可用的 PDF 解析依赖，无法读取 .pdf 文件") from fitz_exc
            raise ValueError("当前环境未安装 PDF 解析依赖，无法读取 .pdf 文件") from exc
        from io import BytesIO

        texts: list[str] = []
        with pdfplumber.open(BytesIO(content)) as pdf:
            for page in pdf.pages:
                texts.append(page.extract_text() or "")
        cleaned = _clean_pdf_text(texts)
        if cleaned:
            return cleaned
        return "\n".join(part for part in texts if part)

    raise ValueError(f"暂不支持的文件类型: {suffix or 'unknown'}")


def split_into_chunks(text: str, chunk_size: int = 900, overlap: int = 120) -> list[str]:
    cleaned = text.replace("\r\n", "\n").strip()
    if not cleaned:
        return []
    structured_chunks = _split_structured_chunks(cleaned, chunk_size, overlap)
    if structured_chunks:
        return structured_chunks
    if len(cleaned) <= chunk_size:
        return [cleaned]
    return _length_chunk_text(cleaned, chunk_size, overlap)


def suggest_question_count_for_chunk(content: str) -> int:
    text = content.strip()
    if not text:
        return 0
    quality = evaluate_chunk_quality(text)
    if quality["score"] < 0.35:
        return 0

    length_score = min(2.2, len(text) / 550)
    sentence_count = max(1, len(re.findall(r"[。！？.!?;\n]", text)))
    structure_bonus = 0.35 if re.search(r"(步骤|流程|注意事项|条件|规则|要求|说明|例如)", text) else 0.0
    density = length_score + min(0.8, sentence_count / 12) + structure_bonus + quality["score"] * 0.35

    if density >= 2.1:
        return 2
    return 1


def allocate_question_counts(
    chunks: Iterable[RagDatasetChunk],
    mode: str,
    requested_count: int | None,
) -> tuple[dict[int, int], int]:
    chunk_list = [chunk for chunk in chunks if (chunk.content or "").strip()]
    if not chunk_list:
        return {}, 0

    weights = {
        chunk.id: float(suggest_question_count_for_chunk(chunk.content))
        for chunk in chunk_list
    }
    allocations = {chunk.id: 0 for chunk in chunk_list}
    eligible_chunks = [chunk for chunk in chunk_list if weights[chunk.id] > 0]
    if not eligible_chunks:
        return allocations, 0

    suggested_total = max(1, int(round(sum(weights[chunk.id] for chunk in eligible_chunks))))
    target_total = suggested_total if mode == "auto" or not requested_count else max(1, requested_count)
    remaining = target_total
    ordered = sorted(eligible_chunks, key=lambda item: (weights[item.id], item.char_count), reverse=True)

    for chunk in ordered:
        if remaining <= 0:
            break
        allocations[chunk.id] += 1
        remaining -= 1

    if remaining > 0:
        total_weight = sum(weights.values()) or 1.0
        fractional: list[tuple[float, int]] = []
        for chunk in ordered:
            share = target_total * weights[chunk.id] / total_weight
            extra = max(0.0, share - allocations[chunk.id])
            fractional.append((extra, chunk.id))
        fractional.sort(reverse=True)
        idx = 0
        while remaining > 0 and fractional:
            chunk_id = fractional[idx % len(fractional)][1]
            allocations[chunk_id] += 1
            remaining -= 1
            idx += 1

    return allocations, suggested_total


def _extract_json_object(text: str) -> dict[str, Any] | None:
    if not text:
        return None
    fenced = re.search(r"```json\s*(.*?)```", text, re.DOTALL | re.IGNORECASE)
    if fenced:
        text = fenced.group(1).strip()
    try:
        loaded = json.loads(text)
        return loaded if isinstance(loaded, dict) else None
    except json.JSONDecodeError:
        pass
    start = text.find("{")
    end = text.rfind("}")
    if start >= 0 and end > start:
        try:
            loaded = json.loads(text[start : end + 1])
            return loaded if isinstance(loaded, dict) else None
        except json.JSONDecodeError:
            pass
    return None


def _coerce_float(value: Any) -> float | None:
    try:
        if value is None or value == "":
            return None
        return float(value)
    except (TypeError, ValueError):
        return None


class OpenAICompatibleClient:
    def __init__(self, llm_config: LLMConfig):
        from openai import AsyncOpenAI

        self.client = AsyncOpenAI(
            base_url=llm_config.api_base_url,
            api_key=llm_config.api_key,
            timeout=180,
            max_retries=2,
        )
        self.model = llm_config.model_name
        self.temperature = llm_config.temperature if llm_config.temperature is not None else 0.2
        self.max_tokens = min(int(llm_config.max_tokens or 2048), 2000)

    async def chat(self, messages: list[dict[str, str]], as_json: bool = False) -> str:
        try:
            response = await self.client.chat.completions.create(
                model=self.model,
                messages=messages,
                temperature=self.temperature,
                max_tokens=self.max_tokens,
                response_format={"type": "json_object"} if as_json else None,
            )
        except Exception:
            response = await self.client.chat.completions.create(
                model=self.model,
                messages=messages,
                temperature=self.temperature,
                max_tokens=self.max_tokens,
            )
        return response.choices[0].message.content or ""


async def generate_chunk_samples(
    question_client: OpenAICompatibleClient,
    chunk: RagDatasetChunk,
    filename: str,
    question_count: int,
) -> list[dict[str, str]]:
    request_count = question_count + min(3, max(1, question_count))
    prompt = QUESTION_GENERATION_PROMPT.format(
        filename=filename,
        chunk_key=chunk.chunk_key,
        content=chunk.content[:12000],
        n=request_count,
    )
    content = await question_client.chat([{"role": "user", "content": prompt}], as_json=True)
    payload = _extract_json_object(content)
    items = payload.get("items") if payload else None
    if not isinstance(items, list):
        raise ValueError("问题生成模型未返回 items 数组")

    samples: list[dict[str, str]] = []
    for item in items:
        if not isinstance(item, dict):
            continue
        question = str(item.get("question") or "").strip()
        reference = str(item.get("reference") or "").strip()
        if question and reference:
            samples.append({"question": question, "reference": reference})
    if not samples:
        raise ValueError("问题生成结果为空")
    validated = await validate_generated_samples(question_client, chunk.content, samples)
    if not validated:
        raise ValueError("候选问题经校验后全部无效")
    return validated[:question_count]


async def validate_generated_samples(
    question_client: OpenAICompatibleClient,
    chunk_content: str,
    samples: list[dict[str, str]],
) -> list[dict[str, str]]:
    deduped: list[dict[str, str]] = []
    seen_questions: set[str] = set()
    for item in samples:
        normalized_question = re.sub(r"\s+", "", (item.get("question") or "")).lower()
        if not normalized_question or normalized_question in seen_questions:
            continue
        if "文档" in normalized_question and ("提到了什么" in normalized_question or "说了什么" in normalized_question):
            continue
        seen_questions.add(normalized_question)
        deduped.append(item)

    if not deduped:
        return []

    prompt = QUESTION_VALIDATION_PROMPT.format(
        content=chunk_content[:12000],
        items_json=json.dumps(deduped, ensure_ascii=False, separators=(",", ":")),
    )
    content = await question_client.chat([{"role": "user", "content": prompt}], as_json=True)
    payload = _extract_json_object(content)
    verdict_items = payload.get("items") if payload else None
    if not isinstance(verdict_items, list):
        return deduped

    verdict_by_question: dict[str, tuple[bool, float]] = {}
    for item in verdict_items:
        if not isinstance(item, dict):
            continue
        question = str(item.get("question") or "").strip()
        keep = bool(item.get("keep"))
        score = _coerce_float(item.get("score")) or 0.0
        if question:
            verdict_by_question[question] = (keep, score)

    kept = [
        (item, verdict_by_question.get(item["question"], (True, 0.5))[1])
        for item in deduped
        if verdict_by_question.get(item["question"], (True, 0.5))[0]
    ]
    kept.sort(key=lambda pair: pair[1], reverse=True)
    return [item for item, _score in kept]


def build_rag_dataset_field_schema() -> list[dict[str, Any]]:
    schema = [
        {"name": "user_input", "type": "text", "required": True, "description": "自动生成的问题"},
        {"name": "reference", "type": "text", "required": True, "description": "基于知识库分片生成的标准答案"},
        {
            "name": "reference_context_ids",
            "type": "text_list",
            "required": True,
            "description": "标准证据分片 ID，用于 HitRate@K / MRR 等检索指标",
        },
        {
            "name": "source_chunk_ids",
            "type": "text_list",
            "required": True,
            "description": "该题来源的知识库分片 ID，便于回溯和 chunk 预览",
        },
        {
            "name": "source_document_names",
            "type": "text_list",
            "required": False,
            "description": "该题对应的源文档名称",
        },
        {
            "name": "generation_meta",
            "type": "json",
            "required": False,
            "description": "生成任务、文档、分片等元信息",
        },
    ]
    return schema


def summarize_supported_metrics() -> tuple[list[str], list[str], list[str]]:
    supported = ["factual_correctness", "answer_completeness"]
    unsupported = ["answer_relevancy", "faithfulness", "context_recall", "context_precision", "contextual_relevancy", "retrieval_hit_rate", "retrieval_mrr"]
    notes = [
        "文档生成数据源只产出 user_input、reference 和标准证据分片，不再调用被测接口生成 response。",
        "需要评测 response 或 retrieved_contexts 时，请在评测执行中选择被测接口，由执行链路实时写入相关字段。",
    ]
    return supported, unsupported, notes


def build_dataset_row_payload(sample: RagDatasetSample) -> dict[str, Any]:
    payload: dict[str, Any] = {
        "user_input": sample.question,
        "reference": sample.reference,
        "reference_context_ids": sample.reference_context_ids or [],
        "source_chunk_ids": sample.source_chunk_ids or [],
        "source_document_names": [sample.document.filename] if sample.document else [],
        "generation_meta": {
            "sample_id": sample.id,
            "job_id": sample.job_id,
            "chunk_id": sample.chunk_id,
        },
    }
    return payload


def _load_job(db: Session, job_id: int) -> RagDatasetJob | None:
    return (
        db.query(RagDatasetJob)
        .options(
            joinedload(RagDatasetJob.documents).joinedload(RagDatasetDocument.chunks),
            joinedload(RagDatasetJob.samples).joinedload(RagDatasetSample.document),
            joinedload(RagDatasetJob.question_llm_config),
            joinedload(RagDatasetJob.dataset),
        )
        .filter(RagDatasetJob.id == job_id)
        .first()
    )


def _log(job: RagDatasetJob, message: str) -> None:
    ts = datetime.now().strftime("%H:%M:%S")
    job.logs = (job.logs or "") + f"[{ts}] {message}\n"
    logger.info("[RAG Builder %s] %s", job.id, message)


def _recount_job(job: RagDatasetJob) -> None:
    selected_samples = [sample for sample in job.samples if sample.selected]
    job.total_samples = len(selected_samples)
    job.completed_samples = sum(1 for sample in selected_samples if sample.status == "completed")
    job.failed_samples = sum(1 for sample in selected_samples if sample.status == "failed")
    job.total_documents = len(job.documents)
    job.total_chunks = sum(len(document.chunks) for document in job.documents)


def _coerce_utc(dt: datetime | None) -> datetime | None:
    if dt is None:
        return None
    if dt.tzinfo is None:
        return dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(timezone.utc)


def _remaining_timeout(deadline: datetime, cap_seconds: float) -> float:
    remaining = (deadline - datetime.now(timezone.utc)).total_seconds()
    if remaining <= 0:
        raise TimeoutError(f"任务执行超时（超过 {RAG_JOB_TIMEOUT_SECONDS // 60} 分钟）")
    return max(0.001, min(cap_seconds, remaining))


def sync_generated_dataset(db: Session, job: RagDatasetJob) -> None:
    # Architecture note:
    # This is the current persistence boundary for generated test data. The RAG
    # job keeps document/chunk/sample execution metadata, but the reusable output
    # consumed by evaluation is always a normal Dataset with DatasetRow records.
    # When a unified DataSource/DatasetGenerator layer is introduced, document
    # generation, API sampling, log replay, and manual-label imports should all
    # converge on this same Dataset/DatasetRow contract instead of making the
    # evaluation flow depend on source-specific job tables.
    supported_metrics, unsupported_metrics, notes = summarize_supported_metrics()
    completed_samples = [
        sample for sample in job.samples if sample.selected and sample.status == "completed"
    ]
    completed_samples.sort(key=lambda item: item.id)

    dataset_name = job.name.strip() or f"RAG 自动数据集 #{job.id}"
    description_lines = [
        "自动生成的 RAG 评测数据集。",
        f"支持指标: {', '.join(supported_metrics)}" if supported_metrics else "支持指标: 无",
    ]
    if unsupported_metrics:
        description_lines.append(f"当前不支持的指标: {', '.join(unsupported_metrics)}")
    description_lines.extend(notes)
    description = "\n".join(description_lines)

    if job.dataset is None:
        dataset = Dataset(
            name=dataset_name,
            description=description,
            sample_type="single_turn",
            field_schema=build_rag_dataset_field_schema(),
            row_count=0,
        )
        db.add(dataset)
        db.flush()
        job.dataset = dataset
        job.dataset_id = dataset.id
    else:
        dataset = job.dataset
        dataset.name = dataset_name
        dataset.description = description
        dataset.sample_type = "single_turn"
        dataset.field_schema = build_rag_dataset_field_schema()

    next_row_index = (
        db.query(DatasetRow.row_index)
        .filter(DatasetRow.dataset_id == dataset.id)
        .order_by(DatasetRow.row_index.desc())
        .first()
    )
    next_row_index = (next_row_index[0] + 1) if next_row_index else 0

    for sample in completed_samples:
        payload = build_dataset_row_payload(sample)
        if sample.dataset_row_id:
            row = db.query(DatasetRow).filter(DatasetRow.id == sample.dataset_row_id).first()
            if row is not None:
                row.data = payload
                continue
        row = DatasetRow(
            dataset_id=dataset.id,
            row_index=next_row_index,
            data=payload,
        )
        db.add(row)
        db.flush()
        sample.dataset_row_id = row.id
        next_row_index += 1

    dataset.row_count = (
        db.query(DatasetRow).filter(DatasetRow.dataset_id == dataset.id).count()
    )


def recover_stale_rag_dataset_job(
    db: Session,
    job_id: int,
    timeout_seconds: int = RAG_JOB_TIMEOUT_SECONDS,
) -> RagDatasetJob | None:
    job = _load_job(db, job_id)
    if job is None or job.status != "running":
        return job

    reference_time = _coerce_utc(job.started_at) or _coerce_utc(job.updated_at) or _coerce_utc(job.created_at)
    if reference_time is None:
        return job

    age_seconds = (datetime.now(timezone.utc) - reference_time).total_seconds()
    if age_seconds < timeout_seconds:
        return job

    timeout_message = f"任务执行超时（超过 {max(1, timeout_seconds // 60)} 分钟），已自动标记为失败"
    for document in job.documents or []:
        for chunk in document.chunks or []:
            if chunk.generation_status == "running":
                chunk.generation_status = "failed"
                chunk.generation_error = timeout_message
    for sample in job.samples or []:
        if sample.status == "running":
            sample.status = "failed"
            sample.error_message = timeout_message
            sample.retry_count = (sample.retry_count or 0) + 1

    if any(sample.selected and sample.status == "completed" for sample in job.samples) or job.dataset is not None:
        sync_generated_dataset(db, job)
    _recount_job(job)
    job.status = "failed" if job.completed_samples == 0 else "partial"
    job.error_message = timeout_message
    job.finished_at = datetime.now(timezone.utc)
    _log(job, f"✗ 自动回收陈旧运行任务：{timeout_message}")
    db.commit()
    db.refresh(job)
    return _load_job(db, job_id)


async def run_rag_dataset_job(
    job_id: int,
    session_factory,
    scope: str = "full",
) -> None:
    db: Session = session_factory()
    try:
        deadline = datetime.now(timezone.utc) + timedelta(seconds=RAG_JOB_TIMEOUT_SECONDS)
        job = _load_job(db, job_id)
        if job is None:
            logger.error("RAG dataset job %s not found", job_id)
            return

        if job.status == "running":
            _log(job, "已有相同任务在运行，忽略本次触发。")
            db.commit()
            return

        if not job.documents:
            job.status = "failed"
            job.error_message = "请先上传知识库文档"
            _log(job, "✗ 没有可用文档")
            db.commit()
            return

        question_client = OpenAICompatibleClient(job.question_llm_config)

        job.status = "running"
        job.error_message = None
        job.started_at = datetime.now(timezone.utc)
        job.finished_at = None
        _log(job, f"========== 开始 RAG 数据集生成（scope={scope}） ==========")

        all_chunks = [chunk for document in job.documents for chunk in document.chunks]
        allocations, suggested_total = allocate_question_counts(
            all_chunks, job.question_count_mode, job.requested_question_count
        )
        job.suggested_question_count = suggested_total
        for chunk in all_chunks:
            chunk.suggested_question_count = suggest_question_count_for_chunk(chunk.content)
            chunk.allocated_question_count = allocations.get(chunk.id, 0)
        db.commit()

        for document in job.documents:
            for chunk in document.chunks:
                allocated_count = chunk.allocated_question_count or 0
                existing_samples = [
                    sample for sample in job.samples if sample.chunk_id == chunk.id and sample.selected
                ]
                should_generate = False
                if scope == "full":
                    should_generate = allocated_count > 0 and not existing_samples
                elif scope == "retry_failed":
                    should_generate = allocated_count > 0 and (
                        chunk.generation_status == "failed" or not existing_samples
                    )

                if not should_generate:
                    continue

                chunk.generation_status = "running"
                chunk.generation_error = None
                db.commit()
                _log(job, f"生成问题：{document.filename} / {chunk.chunk_key}（{allocated_count} 题）")
                try:
                    generated = await asyncio.wait_for(
                        generate_chunk_samples(
                            question_client,
                            chunk,
                            document.filename,
                            allocated_count,
                        ),
                        timeout=_remaining_timeout(deadline, RAG_QUESTION_GENERATION_TIMEOUT_SECONDS),
                    )
                    if not generated:
                        raise ValueError("未生成任何问题")
                    for item in generated:
                        sample = RagDatasetSample(
                            job_id=job.id,
                            document_id=document.id,
                            chunk_id=chunk.id,
                            question=item["question"],
                            reference=item["reference"],
                            reference_context_ids=[chunk.chunk_key],
                            source_chunk_ids=[chunk.chunk_key],
                            status="completed",
                        )
                        db.add(sample)
                    chunk.generation_status = "completed"
                    _log(job, f"✓ 已生成 {len(generated)} 条候选问题")
                except TimeoutError as exc:
                    chunk.generation_status = "failed"
                    chunk.generation_error = str(exc)[:1000]
                    _log(job, f"✗ 问题生成超时：{exc}")
                    db.commit()
                    raise
                except Exception as exc:
                    chunk.generation_status = "failed"
                    chunk.generation_error = str(exc)[:1000]
                    _log(job, f"✗ 问题生成失败：{exc}")
                db.commit()

        db.expire_all()
        job = _load_job(db, job_id)
        if job is None:
            return

        sync_generated_dataset(db, job)
        _recount_job(job)

        if job.total_samples == 0:
            job.status = "failed"
            job.error_message = "没有成功生成任何样本"
            _log(job, "✗ 未生成可用样本")
        elif job.completed_samples > 0 and job.failed_samples == 0:
            job.status = "completed"
            _log(job, f"✓ 全部完成，共 {job.completed_samples} 条样本")
        elif job.completed_samples > 0:
            job.status = "partial"
            _log(job, f"⚠ 部分完成：成功 {job.completed_samples}，失败 {job.failed_samples}")
        else:
            job.status = "failed"
            job.error_message = "所有样本都生成失败"
            _log(job, "✗ 所有样本均失败")
        job.finished_at = datetime.now(timezone.utc)
        _log(job, "========== RAG 数据集生成结束 ==========")
        db.commit()
    except Exception as exc:
        logger.exception("RAG dataset job %s failed", job_id)
        db.rollback()
        job = _load_job(db, job_id)
        if job is not None:
            error_message = str(exc).strip()
            if isinstance(exc, TimeoutError) and not error_message:
                error_message = f"任务执行超时（超过 {RAG_JOB_TIMEOUT_SECONDS // 60} 分钟）"
            job.status = "failed"
            job.error_message = error_message[:2000]
            job.finished_at = datetime.now(timezone.utc)
            _log(job, f"✗✗✗ 任务异常终止：{error_message}")
            db.commit()
    finally:
        db.close()


def persist_uploaded_document(
    db: Session,
    job: RagDatasetJob,
    filename: str,
    raw_bytes: bytes,
    raw_text: str,
) -> RagDatasetDocument:
    upload_root = Path(settings.UPLOAD_DIR).resolve() / "rag_jobs" / str(job.id)
    upload_root.mkdir(parents=True, exist_ok=True)
    stored_name = f"{uuid.uuid4().hex[:8]}_{Path(filename).name}"
    stored_path = upload_root / stored_name
    stored_path.write_bytes(raw_bytes)

    content_hash = hashlib.sha256(raw_bytes).hexdigest()
    document = RagDatasetDocument(
        job_id=job.id,
        filename=filename,
        file_path=str(stored_path),
        content_hash=content_hash,
        raw_content=raw_text,
        char_count=len(raw_text),
    )
    db.add(document)
    db.flush()

    chunks = split_into_chunks(raw_text)
    for index, content in enumerate(chunks):
        chunk = RagDatasetChunk(
            document_id=document.id,
            chunk_index=index,
            chunk_key=f"doc-{document.id}-chunk-{index}",
            content=content,
            char_count=len(content),
            suggested_question_count=suggest_question_count_for_chunk(content),
            allocated_question_count=0,
            generation_status="pending",
        )
        db.add(chunk)

    document.chunk_count = len(chunks)
    return document
